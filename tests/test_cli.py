"""The command line, and one pass through the whole pipeline without an API key.

`main.run` takes an injected client, so the end-to-end test below drives a real
PPTX through ingestion, extraction, normalisation, valuation, validation and both
outputs with a stub standing in for the model. That is the only test that proves
the pieces fit together; everything else tests them apart.
"""

from __future__ import annotations

import argparse
import json

import pytest

from app import config, main
from app.extraction import company_extractor
from app.output import json_export


# --- A model that answers from a script -------------------------------------


class StubClient:
    """Returns a fixed extraction response, shaped exactly as the schema requires."""

    def __init__(self, payload: dict):
        self.payload = payload
        self.calls = 0

    class _Stream:
        def __init__(self, text: str):
            self._text = text

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def get_final_message(self):
            block = type("Block", (), {"type": "text", "text": self._text})()
            return type("Message", (), {"stop_reason": "end_turn", "content": [block]})()

    @property
    def messages(self):
        outer = self

        class Messages:
            def stream(self, **kwargs):
                outer.calls += 1
                return StubClient._Stream(json.dumps(outer.payload))

        return Messages()


@pytest.fixture
def extraction_payload() -> dict:
    def fact(field, text, slide):
        return {"field": field, "text": text, "slide": slide, "quote": text}

    return {
        "facts": [
            fact("company.name", "Northwind Analytics", 1),
            fact("company.sector", "B2B SaaS", 1),
            fact("company.geography", "United States", 1),
            fact("company.stage", "Series A", 1),
            fact("financing.raise_amount", "$4M", 1),
            fact("financing.pre_money", "$18M", 1),
            fact("financing.instrument", "Priced equity", 1),
            fact("metrics.arr", "$2.4M", 2),
            fact("metrics.revenue_label", "2026 actual", 2),
            fact("metrics.nrr", "112%", 2),
            fact("metrics.growth_rate", "94%", 2),
            fact("metrics.gross_margin", "76%", 2),
            fact("metrics.customer_count", "38", 2),
        ],
        "scorecard": {
            factor: {"evidence": "present", "score": 1, "rationale": "Shown on slide 2."}
            for factor in company_extractor.SCORECARD_FACTORS
        },
        "contradictions": [],
    }


@pytest.fixture
def deck_file(tmp_path):
    """A real PPTX, so the ingestion layer is exercised rather than mocked."""
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    layout = presentation.slide_layouts[5]

    cover = presentation.slides.add_slide(layout)
    cover.shapes.title.text = "Northwind Analytics"
    frame = cover.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)).text_frame
    frame.text = "Series A | B2B SaaS | United States"
    frame.add_paragraph().text = "Raising $4M at $18M pre-money"

    traction = presentation.slides.add_slide(layout)
    traction.shapes.title.text = "Traction"
    table = traction.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for row, (label, value) in enumerate(
        [("Metric", "Value"), ("ARR", "$2.4M"), ("NRR", "112%")]
    ):
        table.cell(row, 0).text = label
        table.cell(row, 1).text = value
    traction.notes_slide.notes_text_frame.text = "Gross margin 76%, growth 94% YoY."

    path = tmp_path / "northwind.pptx"
    presentation.save(str(path))
    return path


# --- End to end -------------------------------------------------------------


class TestPipeline:
    def test_a_deck_becomes_a_valuation_without_an_api_key(
        self, deck_file, extraction_payload, tmp_path
    ):
        client = StubClient(extraction_payload)
        profile, analysis = main.run(
            deck_file,
            client=client,
            interpret=False,
            progress=main.Progress(echo=False),
        )

        assert client.calls == 1, "the deck should be read exactly once"
        assert profile.identity.name == "Northwind Analytics"
        assert profile.slide_count == 2
        assert profile.stage == config.SERIES_A
        assert profile.metrics.arr == pytest.approx(2_400_000)
        assert profile.metrics.nrr == pytest.approx(1.12)
        assert analysis.range.low <= analysis.range.base <= analysis.range.high
        assert analysis.requested_valuation == pytest.approx(18_000_000)

    def test_the_run_produces_both_artefacts(self, deck_file, extraction_payload, tmp_path):
        from app.reporting import pdf_report

        client = StubClient(extraction_payload)
        profile, analysis = main.run(
            deck_file, client=client, interpret=False, progress=main.Progress(echo=False)
        )
        result = pdf_report.render(profile, analysis, tmp_path / "out.pdf")
        record = json_export.build(profile, analysis)
        path = json_export.write(record, tmp_path / "out.json")

        assert result["pages"] == 1
        assert result["path"].stat().st_size > 5_000
        assert json_export.verify_schema(json.loads(path.read_text(encoding="utf-8"))) == []

    def test_user_input_beats_the_deck_end_to_end(
        self, deck_file, extraction_payload, tmp_path
    ):
        client = StubClient(extraction_payload)
        profile, _ = main.run(
            deck_file,
            overrides={"growth_rate": "60%", "requested_pre_money": 30_000_000},
            client=client,
            interpret=False,
            progress=main.Progress(echo=False),
        )
        assert profile.metrics.growth_rate == pytest.approx(0.60)
        assert profile.ledger.get("growth_rate").source_type == config.USER
        assert profile.round.requested_pre_money == pytest.approx(30_000_000)

    def test_interpretation_failure_leaves_the_numbers_intact(
        self, deck_file, extraction_payload
    ):
        """The written pass is a convenience; losing it must not lose the analysis."""
        client = StubClient(extraction_payload)
        deterministic, _ = main.run(
            deck_file, client=client, interpret=False, progress=main.Progress(echo=False)
        )
        from app.valuation import valuation_engine

        baseline = valuation_engine.run(deterministic)

        # A client that refuses every call stands in for a failed interpretation.
        class Broken(StubClient):
            @property
            def messages(self):
                raise RuntimeError("no model available")

        from app import narrative

        analysis = narrative.refine(deterministic, baseline, client=Broken({}), enabled=True)
        assert analysis.range.base == pytest.approx(baseline.range.base)
        assert analysis.key_milestone


# --- Argument handling ------------------------------------------------------


def namespace(**overrides) -> argparse.Namespace:
    defaults = dict(
        inputs=None,
        company=None,
        user_name=None,
        user_email=None,
        raise_amount=None,
        pre_money=None,
        consent_analysis=False,
        consent_storage=False,
        consent_contact=False,
    )
    defaults.update(overrides)
    return argparse.Namespace(**defaults)


class TestInputMerging:
    def test_no_inputs_is_an_empty_mapping(self):
        assert main._load_inputs(namespace()) == {}

    def test_the_command_line_beats_the_input_file(self, tmp_path):
        path = tmp_path / "inputs.json"
        path.write_text(json.dumps({"company_name": "From File", "arr": 1_000_000}))
        merged = main._load_inputs(namespace(inputs=str(path), company="From Flag"))
        assert merged["company_name"] == "From Flag"
        assert merged["arr"] == 1_000_000

    def test_a_missing_input_file_is_an_explicit_error(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            main._load_inputs(namespace(inputs=str(tmp_path / "absent.json")))

    def test_a_non_object_input_file_is_rejected(self, tmp_path):
        path = tmp_path / "inputs.json"
        path.write_text(json.dumps([1, 2, 3]))
        with pytest.raises(ValueError, match="JSON object"):
            main._load_inputs(namespace(inputs=str(path)))

    def test_consent_flags_are_absent_unless_passed(self):
        merged = main._load_inputs(namespace())
        assert "analysis_consent" not in merged

    def test_a_consent_flag_sets_a_literal_true(self):
        merged = main._load_inputs(namespace(consent_storage=True))
        assert merged["data_storage_consent"] is True
        assert "analysis_consent" not in merged

    def test_the_shipped_template_uses_only_recognised_keys(self):
        from app.extraction import profile_builder

        path = config.PROJECT_ROOT / "templates" / "user_inputs.json"
        supplied = json.loads(path.read_text(encoding="utf-8"))
        unknown = [key for key in profile_builder.unknown_keys(supplied) if key != "_comment"]
        assert unknown == []

    def test_the_template_grants_no_consent(self):
        path = config.PROJECT_ROOT / "templates" / "user_inputs.json"
        supplied = json.loads(path.read_text(encoding="utf-8"))
        for flag in ("analysis_consent", "data_storage_consent", "contact_consent"):
            assert supplied[flag] is False


class TestOutputPaths:
    def test_the_pdf_and_json_share_a_stem(self, saas):
        pdf_path, json_path = main._paths(saas, None)
        assert pdf_path.stem == json_path.stem
        assert pdf_path.suffix == ".pdf" and json_path.suffix == ".json"

    def test_the_default_name_is_a_slug_of_the_company(self, saas):
        pdf_path, _ = main._paths(saas, None)
        assert pdf_path.name == "northwind-analytics_valuation.pdf"

    def test_an_override_without_a_suffix_still_becomes_a_pdf(self, saas, tmp_path):
        pdf_path, json_path = main._paths(saas, str(tmp_path / "report"))
        assert pdf_path.name == "report.pdf"
        assert json_path.name == "report.json"

    def test_a_company_with_no_usable_name_still_gets_a_path(self, incomplete):
        incomplete.identity.name = "!!!"
        pdf_path, _ = main._paths(incomplete, None)
        assert pdf_path.name == "company_valuation.pdf"


class TestArgumentParser:
    def test_the_deck_is_the_only_required_argument(self):
        args = main._parser().parse_args(["deck.pdf"])
        assert args.deck == "deck.pdf"
        assert args.comparables is None
        assert args.consent_analysis is False

    def test_the_documented_flags_all_parse(self):
        args = main._parser().parse_args(
            [
                "deck.pptx",
                "--inputs", "in.json",
                "--comparables", "comps.csv",
                "--output", "out.pdf",
                "--company", "Northwind",
                "--raise-amount", "4000000",
                "--pre-money", "18000000",
                "--consent-analysis",
                "--no-interpretation",
                "--date", "2026-09-03",
            ]
        )
        assert args.raise_amount == pytest.approx(4_000_000)
        assert args.pre_money == pytest.approx(18_000_000)
        assert args.consent_analysis is True
        assert args.no_interpretation is True
        assert args.date == "2026-09-03"
