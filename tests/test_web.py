"""The Flask front end: request handling, job lifecycle, and the rules it must not relax.

The web layer is a wrapper — the numbers a browser gets are the numbers the CLI
would have produced from the same inputs, and nothing analytical is
re-implemented here. So these tests concentrate on what the wrapper itself owns:
what it accepts, what it refuses, how a job moves from queued to done, and
whether the guarantees the standard makes survive the trip through a form.

The pipeline is stubbed throughout, so the suite still runs offline and spends
no tokens.
"""

from __future__ import annotations

import io
import json
import threading
import time

import pytest

flask = pytest.importorskip("flask")

from app import config  # noqa: E402
from web import server  # noqa: E402


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test")
    monkeypatch.delenv("APP_PASSWORD", raising=False)
    server._jobs.clear()
    server.flask_app.config["TESTING"] = True
    return server.flask_app.test_client()


@pytest.fixture
def deck_bytes() -> bytes:
    """A tiny real PPTX, so the upload path is exercised rather than mocked."""
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Placeholder Deck"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text_frame.text = (
        "Raising $4M at $18M pre-money"
    )
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def upload(deck_bytes: bytes, name: str = "deck.pptx", **fields):
    data = {"deck": (io.BytesIO(deck_bytes), name)}
    data.update(fields)
    return data


@pytest.fixture
def stub_pipeline(monkeypatch, saas):
    """Replace the pipeline with one that finishes immediately."""
    from app.utils import validation
    from app.valuation import valuation_engine

    analysis = valuation_engine.run(saas)
    analysis.issues = validation.run(saas, analysis)
    captured: dict = {}

    def fake_run(deck_path, overrides=None, comparables=None, interpret=True, progress=None, **_):
        captured["overrides"] = overrides
        captured["comparables"] = comparables
        captured["interpret"] = interpret
        if progress is not None:
            progress.step(1, "Reading deck")
            progress.detail("2 slides")
        saas.user.name = (overrides or {}).get("user_name", "")
        saas.user.email = (overrides or {}).get("user_email", "")
        saas.user.consent.analysis_consent = (overrides or {}).get("analysis_consent") is True
        saas.user.consent.data_storage_consent = (
            (overrides or {}).get("data_storage_consent") is True
        )
        saas.user.consent.contact_consent = (overrides or {}).get("contact_consent") is True
        saas.user.consent.timestamp = None
        saas.user.consent.stamp()
        return saas, analysis

    monkeypatch.setattr(server.pipeline, "run", fake_run)
    return captured


def wait_for(client, job_id: str, timeout: float = 15.0) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        payload = client.get(f"/jobs/{job_id}").get_json()
        if payload["status"] in ("done", "error"):
            return payload
        time.sleep(0.05)
    raise AssertionError("job did not finish")


# --- Pages and probes -------------------------------------------------------


class TestPages:
    def test_the_index_renders(self, client):
        response = client.get("/")
        assert response.status_code == 200
        body = response.get_data(as_text=True)
        assert "Valuation Analysis" in body
        assert config.LLM_MODEL in body

    def test_the_page_offers_every_control_the_server_reads(self, client):
        """A rename in the markup must not silently orphan a form field."""
        body = client.get("/").get_data(as_text=True)
        for name in (
            "deck",
            "comparables",
            "inputs",
            "company",
            "sector",
            "geography",
            "raise_amount",
            "pre_money",
            "valuation_cap",
            "arr",
            "growth_rate",
            "gross_margin",
            "nrr",
            "revenue_basis",
            "user_name",
            "user_email",
            "consent_analysis",
            "consent_storage",
            "consent_contact",
            "no_interpretation",
        ):
            assert f'name="{name}"' in body, name

    def test_the_page_advertises_only_supported_deck_formats(self, client):
        """Offering a format the ingestion layer rejects is a promise the app breaks."""
        body = client.get("/").get_data(as_text=True)
        assert ".docx" not in body
        assert 'accept=".pdf,.pptx,.ppt"' in body

    def test_the_brand_rule_is_not_clipped_by_its_own_card(self, client):
        """The tri-colour rule is the brand signature; a negative top hides it entirely.

        The card sets `overflow:hidden`, so a rule positioned above the content
        box is clipped away completely and the card renders with no signature at
        all — which is exactly what happened, and is invisible in any test that
        only checks the markup is present.
        """
        body = client.get("/").get_data(as_text=True)
        rule = body.split(".card::after{", 1)[1].split("}", 1)[0]
        assert "top: 0;" in rule or "top:0;" in rule
        assert "top:-" not in rule.replace(" ", "")

    def test_the_index_lists_every_pipeline_stage(self, client):
        body = client.get("/").get_data(as_text=True)
        for label in server.pipeline.STEP_LABELS:
            assert label in body

    def test_the_index_states_the_no_comparables_position(self, client):
        body = client.get("/").get_data(as_text=True)
        assert "never fetched and never invented" in body
        assert "ANALYST ASSUMPTION" in body

    def test_healthz_reports_readiness_without_spending_tokens(self, client):
        payload = client.get("/healthz").get_json()
        assert payload["ok"] is True
        assert payload["api_key_configured"] is True
        assert payload["model"] == config.LLM_MODEL
        assert "api" not in payload

    def test_healthz_reports_a_missing_key(self, client, monkeypatch):
        monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
        monkeypatch.delenv("ANTHROPIC_AUTH_TOKEN", raising=False)
        payload = client.get("/healthz").get_json()
        assert payload["ok"] is False
        assert payload["api_key_configured"] is False

    def test_the_favicon_is_served_at_the_root(self, client):
        assert client.get("/favicon.ico").status_code == 200

    def test_the_template_is_served_for_reference(self, client):
        payload = client.get("/template").get_json()
        assert payload["template_name"].startswith("TEN Capital")
        assert "sections" in payload


# --- What the upload endpoint refuses ---------------------------------------


class TestRejections:
    def test_no_deck_is_refused(self, client):
        response = client.post("/analyze", data={})
        assert response.status_code == 400
        assert "No pitch deck" in response.get_json()["error"]

    def test_an_unsupported_deck_type_is_refused_by_name(self, client, deck_bytes):
        response = client.post(
            "/analyze", data=upload(deck_bytes, "deck.docx"), content_type="multipart/form-data"
        )
        assert response.status_code == 400
        assert ".docx" in response.get_json()["error"]

    def test_a_missing_api_key_is_refused_with_a_fix(self, client, deck_bytes, monkeypatch):
        monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
        monkeypatch.delenv("ANTHROPIC_AUTH_TOKEN", raising=False)
        response = client.post(
            "/analyze", data=upload(deck_bytes), content_type="multipart/form-data"
        )
        assert response.status_code == 503
        assert "Railway service variables" in response.get_json()["error"]

    def test_a_wrong_password_is_refused(self, client, deck_bytes, monkeypatch):
        monkeypatch.setenv("APP_PASSWORD", "correct-horse")
        response = client.post(
            "/analyze",
            data=upload(deck_bytes, password="wrong"),
            content_type="multipart/form-data",
        )
        assert response.status_code == 401

    def test_the_right_password_is_accepted(self, client, deck_bytes, stub_pipeline, monkeypatch):
        monkeypatch.setenv("APP_PASSWORD", "correct-horse")
        response = client.post(
            "/analyze",
            data=upload(deck_bytes, password="correct-horse"),
            content_type="multipart/form-data",
        )
        assert response.status_code == 202

    def test_an_unsupported_comparables_file_is_refused(self, client, deck_bytes):
        data = upload(deck_bytes)
        data["comparables"] = (io.BytesIO(b"x"), "comps.xlsx")
        response = client.post("/analyze", data=data, content_type="multipart/form-data")
        assert response.status_code == 400
        assert ".xlsx" in response.get_json()["error"]

    def test_a_malformed_inputs_file_is_refused(self, client, deck_bytes):
        data = upload(deck_bytes)
        data["inputs"] = (io.BytesIO(b"{not json"), "inputs.json")
        response = client.post("/analyze", data=data, content_type="multipart/form-data")
        assert response.status_code == 400
        assert "not valid JSON" in response.get_json()["error"]

    def test_an_unrecognised_input_key_is_reported_not_ignored(self, client, deck_bytes):
        data = upload(deck_bytes)
        data["inputs"] = (io.BytesIO(json.dumps({"revenu": 1}).encode()), "inputs.json")
        response = client.post("/analyze", data=data, content_type="multipart/form-data")
        assert response.status_code == 400
        assert "revenu" in response.get_json()["error"]

    def test_the_concurrency_ceiling_is_enforced(self, client, deck_bytes, monkeypatch):
        monkeypatch.setattr(server, "MAX_CONCURRENT_JOBS", 0)
        response = client.post(
            "/analyze", data=upload(deck_bytes), content_type="multipart/form-data"
        )
        assert response.status_code == 429

    def test_an_oversized_upload_returns_json_not_html(self, client, monkeypatch):
        monkeypatch.setitem(server.flask_app.config, "MAX_CONTENT_LENGTH", 64)
        response = client.post(
            "/analyze",
            data=upload(b"x" * 4096),
            content_type="multipart/form-data",
        )
        assert response.status_code == 413
        assert "larger than" in response.get_json()["error"]


# --- The job lifecycle ------------------------------------------------------


class TestJobLifecycle:
    def test_an_upload_returns_a_job_immediately(self, client, deck_bytes, stub_pipeline):
        response = client.post(
            "/analyze", data=upload(deck_bytes), content_type="multipart/form-data"
        )
        assert response.status_code == 202
        job = response.get_json()
        assert job["status"] in ("queued", "running")
        assert job["total"] == server.pipeline.TOTAL_STEPS

    def test_the_job_reaches_done_and_carries_the_headline(
        self, client, deck_bytes, stub_pipeline
    ):
        job = client.post(
            "/analyze", data=upload(deck_bytes), content_type="multipart/form-data"
        ).get_json()
        finished = wait_for(client, job["id"])
        assert finished["status"] == "done"
        summary = finished["summary"]
        for key in ("low", "base", "high", "status", "confidence", "methods", "dilution"):
            assert key in summary
        assert summary["status"] in set(config.STATUS_COLORS)

    def test_the_public_view_never_leaks_file_bytes(self, client, deck_bytes, stub_pipeline):
        job = client.post(
            "/analyze", data=upload(deck_bytes), content_type="multipart/form-data"
        ).get_json()
        finished = wait_for(client, job["id"])
        assert "pdf" not in finished and "json" not in finished

    def test_both_artefacts_download(self, client, deck_bytes, stub_pipeline):
        job = client.post(
            "/analyze", data=upload(deck_bytes), content_type="multipart/form-data"
        ).get_json()
        wait_for(client, job["id"])

        pdf = client.get(f"/jobs/{job['id']}/pdf")
        assert pdf.status_code == 200
        assert pdf.mimetype == "application/pdf"
        assert pdf.data.startswith(b"%PDF")

        sidecar = client.get(f"/jobs/{job['id']}/json")
        assert sidecar.status_code == 200
        record = json.loads(sidecar.data)
        from app.output import json_export

        assert json_export.verify_schema(record) == []

    def test_the_download_name_is_the_company_slug(self, client, deck_bytes, stub_pipeline):
        job = client.post(
            "/analyze", data=upload(deck_bytes), content_type="multipart/form-data"
        ).get_json()
        wait_for(client, job["id"])
        response = client.get(f"/jobs/{job['id']}/pdf?download=1")
        assert "northwind-analytics_valuation.pdf" in response.headers["Content-Disposition"]

    def test_artefacts_are_refused_while_the_job_is_running(self, client, deck_bytes, monkeypatch):
        release = threading.Event()

        def slow_run(*args, **kwargs):
            release.wait(timeout=10)
            raise ValueError("stopped")

        monkeypatch.setattr(server.pipeline, "run", slow_run)
        job = client.post(
            "/analyze", data=upload(deck_bytes), content_type="multipart/form-data"
        ).get_json()
        try:
            assert client.get(f"/jobs/{job['id']}/pdf").status_code == 409
        finally:
            release.set()

    def test_a_failed_run_surfaces_its_reason(self, client, deck_bytes, monkeypatch):
        def broken(*args, **kwargs):
            raise ValueError("the deck yielded no extractable text")

        monkeypatch.setattr(server.pipeline, "run", broken)
        job = client.post(
            "/analyze", data=upload(deck_bytes), content_type="multipart/form-data"
        ).get_json()
        finished = wait_for(client, job["id"])
        assert finished["status"] == "error"
        assert "no extractable text" in finished["error"]
        assert client.get(f"/jobs/{job['id']}/pdf").status_code == 500

    def test_an_unknown_job_is_a_clean_404(self, client):
        assert client.get("/jobs/deadbeef").status_code == 404

    def test_finished_jobs_are_reaped_past_their_ttl(self, client, deck_bytes, stub_pipeline):
        job = client.post(
            "/analyze", data=upload(deck_bytes), content_type="multipart/form-data"
        ).get_json()
        wait_for(client, job["id"])
        server._jobs[job["id"]]["finished_at"] = time.time() - server.JOB_TTL_SECONDS - 1
        server._reap()
        assert client.get(f"/jobs/{job['id']}").status_code == 404


# --- The guarantees the form must not relax ---------------------------------


class TestFormHonoursTheStandard:
    def test_consent_is_off_unless_the_box_is_ticked(self, client, deck_bytes, stub_pipeline):
        client.post(
            "/analyze",
            data=upload(deck_bytes, consent_analysis="0", consent_storage="0"),
            content_type="multipart/form-data",
        )
        overrides = stub_pipeline["overrides"]
        assert "analysis_consent" not in overrides
        assert "data_storage_consent" not in overrides

    def test_a_ticked_box_sets_a_literal_true(self, client, deck_bytes, stub_pipeline):
        client.post(
            "/analyze",
            data=upload(deck_bytes, consent_analysis="1"),
            content_type="multipart/form-data",
        )
        assert stub_pipeline["overrides"]["analysis_consent"] is True

    def test_consent_reaches_the_record_exactly_as_ticked(
        self, client, deck_bytes, stub_pipeline
    ):
        job = client.post(
            "/analyze",
            data=upload(deck_bytes, consent_analysis="1", consent_contact="0"),
            content_type="multipart/form-data",
        ).get_json()
        wait_for(client, job["id"])
        record = json.loads(client.get(f"/jobs/{job['id']}/json").data)
        assert record["consent"]["analysis_consent"] is True
        assert record["consent"]["contact_consent"] is False
        assert record["consent"]["timestamp"]

    def test_an_empty_box_does_not_overwrite_the_deck(self, client, deck_bytes, stub_pipeline):
        """A blank field is silence, not an instruction to blank the deck's figure."""
        client.post(
            "/analyze",
            data=upload(deck_bytes, company="", arr="", growth_rate=""),
            content_type="multipart/form-data",
        )
        overrides = stub_pipeline["overrides"]
        assert "company_name" not in overrides
        assert "arr" not in overrides
        assert "growth_rate" not in overrides

    def test_typed_values_are_parsed_into_the_right_units(
        self, client, deck_bytes, stub_pipeline
    ):
        client.post(
            "/analyze",
            data=upload(
                deck_bytes,
                company="Placeholder Ltd",
                raise_amount="$4M",
                pre_money="$18M",
                arr="$2.4M",
                growth_rate="94%",
                nrr="112%",
            ),
            content_type="multipart/form-data",
        )
        overrides = stub_pipeline["overrides"]
        assert overrides["company_name"] == "Placeholder Ltd"
        assert overrides["raise_amount"] == pytest.approx(4_000_000)
        assert overrides["requested_pre_money"] == pytest.approx(18_000_000)
        assert overrides["arr"] == pytest.approx(2_400_000)
        assert overrides["growth_rate"] == pytest.approx(0.94)
        assert overrides["nrr"] == pytest.approx(1.12)

    def test_a_typed_value_beats_the_uploaded_inputs_file(
        self, client, deck_bytes, stub_pipeline
    ):
        data = upload(deck_bytes, company="From The Form")
        data["inputs"] = (
            io.BytesIO(json.dumps({"company_name": "From The File", "arr": 900_000}).encode()),
            "inputs.json",
        )
        client.post("/analyze", data=data, content_type="multipart/form-data")
        overrides = stub_pipeline["overrides"]
        assert overrides["company_name"] == "From The Form"
        assert overrides["arr"] == pytest.approx(900_000)

    def test_the_revenue_basis_override_is_passed_through(
        self, client, deck_bytes, stub_pipeline
    ):
        client.post(
            "/analyze",
            data=upload(deck_bytes, revenue_basis="projected"),
            content_type="multipart/form-data",
        )
        assert stub_pipeline["overrides"]["revenue_basis"] == "projected"

    def test_an_invalid_revenue_basis_is_dropped_not_passed_on(
        self, client, deck_bytes, stub_pipeline
    ):
        client.post(
            "/analyze",
            data=upload(deck_bytes, revenue_basis="wishful"),
            content_type="multipart/form-data",
        )
        assert "revenue_basis" not in stub_pipeline["overrides"]

    def test_no_comparables_file_means_none_supplied(self, client, deck_bytes, stub_pipeline):
        client.post("/analyze", data=upload(deck_bytes), content_type="multipart/form-data")
        assert stub_pipeline["comparables"] == []

    def test_a_comparables_csv_is_parsed_and_passed_through(
        self, client, deck_bytes, stub_pipeline
    ):
        csv = b"company,stage,pre_money,arr,source\nExample Co,Series A,$24M,$2.6M,PitchBook\n"
        data = upload(deck_bytes)
        data["comparables"] = (io.BytesIO(csv), "comps.csv")
        client.post("/analyze", data=data, content_type="multipart/form-data")
        comparables = stub_pipeline["comparables"]
        assert len(comparables) == 1
        assert comparables[0].company == "Example Co"
        assert comparables[0].pre_money == pytest.approx(24_000_000)

    def test_the_interpretation_can_be_switched_off(self, client, deck_bytes, stub_pipeline):
        client.post(
            "/analyze",
            data=upload(deck_bytes, no_interpretation="1"),
            content_type="multipart/form-data",
        )
        assert stub_pipeline["interpret"] is False


# --- Deployment shape -------------------------------------------------------


class TestDeploymentConfiguration:
    """The registry is process-local, so the deploy has to be a single worker."""

    def test_the_procfile_runs_one_worker(self):
        procfile = (config.PROJECT_ROOT / "Procfile").read_text(encoding="utf-8")
        assert "web.server:app" in procfile
        assert "--workers 1" in procfile
        assert "--threads" in procfile

    def test_railway_json_matches_the_procfile_and_probes_healthz(self):
        railway = json.loads((config.PROJECT_ROOT / "railway.json").read_text(encoding="utf-8"))
        deploy = railway["deploy"]
        assert "--workers 1" in deploy["startCommand"]
        assert "web.server:app" in deploy["startCommand"]
        assert deploy["healthcheckPath"] == "/healthz"

    def test_the_start_timeout_outlasts_a_run(self):
        """A run takes minutes; a short worker timeout would kill it mid-analysis."""
        procfile = (config.PROJECT_ROOT / "Procfile").read_text(encoding="utf-8")
        timeout = int(procfile.split("--timeout")[1].split()[0])
        assert timeout >= 300

    def test_the_build_never_ships_the_env_file(self):
        ignored = (config.PROJECT_ROOT / ".railwayignore").read_text(encoding="utf-8")
        assert ".env" in ignored.splitlines()

    def test_the_env_file_is_gitignored(self):
        ignored = (config.PROJECT_ROOT / ".gitignore").read_text(encoding="utf-8")
        assert ".env" in ignored.splitlines()

    def test_the_example_env_carries_no_key(self):
        example = (config.PROJECT_ROOT / ".env.example").read_text(encoding="utf-8")
        assert "ANTHROPIC_API_KEY=" in example
        assert "sk-ant" not in example

    def test_the_server_dependencies_are_declared(self):
        requirements = (config.PROJECT_ROOT / "requirements.txt").read_text(encoding="utf-8")
        for package in ("Flask", "gunicorn", "resend"):
            assert package in requirements

    def test_requirements_and_pyproject_do_not_drift(self):
        """Railway installs requirements.txt; a local `pip install .` reads pyproject."""
        import re

        def names(text: str) -> set[str]:
            found = re.findall(r'"?([A-Za-z][A-Za-z0-9_.-]*)\s*>=', text)
            return {name.lower() for name in found}

        requirements = (config.PROJECT_ROOT / "requirements.txt").read_text(encoding="utf-8")
        pyproject = (config.PROJECT_ROOT / "pyproject.toml").read_text(encoding="utf-8")
        declared = pyproject.split("[project.optional-dependencies]")[0]
        assert names(requirements) == names(declared)
